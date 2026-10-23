"""Tests for the Qdrant, embeddings, SharePoint MCP and parser tools plus the extended agent endpoint."""
import asyncio
import base64
import contextlib
import io
from types import SimpleNamespace

from fastapi.testclient import TestClient

from src.main import MODELS, app, build_messages, ExecuteRequest
from src.tools.embeddings import EmbeddingsError, EmbeddingsTool
from src.tools.qdrant import QdrantTool, build_filter, point_id, to_chunk
from src.tools.sharepoint_mcp import SharePointMCPError, SharePointMCPTool

client = TestClient(app)


# ---- Agent endpoint ----


def test_execute_json_mock_returns_blocks():
    response = client.post(
        "/agent/execute",
        json={"model": "groq:openai/gpt-oss-20b", "prompt": "Material Status", "live": False,
              "prompt_templates": ["manzil_v2", "manzil_ui_output"], "response_format": "json"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["json"]["blocks"][0]["type"] == "markdown"


def test_execute_unknown_template_is_rejected_in_mock_mode():
    response = client.post("/agent/execute", json={"prompt": "Hi", "live": False, "prompt_templates": ["nope"]})
    assert response.status_code == 400
    assert "manzil_v2" in response.json()["detail"]


def test_execute_rejects_template_path_traversal():
    response = client.post("/agent/execute", json={"prompt": "Hi", "live": False, "prompt_templates": ["../.env"]})
    assert response.status_code == 400


def test_build_messages_orders_templates_history_and_context():
    messages = build_messages(ExecuteRequest(
        prompt="Material Status", prompt_templates=["manzil_v2"], system_prompt="Extra rule.",
        history=[{"role": "user", "content": "earlier"}, {"role": "system", "content": "dropped"}],
        context='[{"id": "c1"}]', response_format="json",
    ))
    assert messages[0]["role"] == "system"
    assert messages[0]["content"].startswith("# PROJECT EXECUTION VALIDATION ASSISTANT")
    assert "Extra rule." in messages[0]["content"] and "JSON object only" in messages[0]["content"]
    assert [m["role"] for m in messages] == ["system", "user", "user"]
    assert messages[-1]["content"].startswith("Retrieved knowledge source context:")


def test_prompts_endpoint_lists_templates():
    assert {"manzil_v2", "manzil_dashboard", "manzil_ui_output"} <= set(client.get("/prompts").json()["templates"])


# ---- Embeddings ----


def test_embeddings_mock_is_deterministic_and_sized():
    body = client.post("/embeddings", json={"texts": ["a", "b", "a"], "live": False}).json()
    assert body["model"] == "gemini:gemini-embedding-2" and body["dimensions"] == 3072
    assert len(body["vectors"]) == 3 and len(body["vectors"][0]) == 3072
    assert body["vectors"][0] == body["vectors"][2] != body["vectors"][1]


def test_embeddings_unknown_model():
    try:
        EmbeddingsTool(MODELS).resolve("openai:nope")
    except EmbeddingsError as err:
        assert "text-embedding-3-small" in str(err)
    else:
        raise AssertionError("expected EmbeddingsError")


# ---- Qdrant ----


def test_qdrant_mock_search_has_citation_metadata():
    body = client.post("/connectors/qdrant/search", json={"collection": "docs", "top_k": 2, "filter": {"project_id": "P9"}, "live": False}).json()
    assert len(body["chunks"]) == 2
    assert body["chunks"][0]["metadata"]["project_id"] == "P9"
    assert body["chunks"][0]["metadata"]["url"].startswith("https://")


def test_qdrant_live_without_url_is_a_clear_error(monkeypatch):
    tool = QdrantTool(url="")
    try:
        asyncio.run(tool.search("docs", [0.1], live=True))
    except Exception as err:
        assert "QDRANT_URL" in str(err)
    else:
        raise AssertionError("expected an error")


def test_qdrant_live_search_flattens_payload():
    class FakeClient:
        async def query_points(self, **kwargs):
            self.kwargs = kwargs
            point = SimpleNamespace(id=7, score=0.81234, payload={"page_content": "hello", "metadata": {"doc_name": "A.pdf"}, "project_id": "P1"})
            return SimpleNamespace(points=[point])

    fake = FakeClient()
    result = asyncio.run(QdrantTool(client=fake).search("docs", [0.1, 0.2], top_k=3, filter={"project_id": "P1", "doc_type": ""}))
    assert result["chunks"] == [{"id": "7", "score": 0.8123, "text": "hello", "metadata": {"doc_name": "A.pdf", "project_id": "P1"}}]
    assert fake.kwargs["limit"] == 3
    assert len(fake.kwargs["query_filter"].must) == 1  # blank doc_type dropped


def test_build_filter_and_point_ids():
    assert build_filter({"a": "", "b": []}) is None
    assert build_filter({"doc_type": ["LOA", "BOQ"]}).must[0].match.any == ["LOA", "BOQ"]
    p = {"vector": [1.0], "payload": {"url": "https://x/a.pdf", "chunk_index": 3}}
    assert point_id(p) == point_id(dict(p)) != point_id({"payload": {"url": "https://x/a.pdf", "chunk_index": 4}})
    assert to_chunk(1, 0.5, {"text": "t", "k": "v"}, "text")["metadata"] == {"k": "v"}


def test_qdrant_upsert_requires_vectors():
    response = client.post("/connectors/qdrant/upsert", json={"collection": "docs", "points": [{"payload": {}}], "live": False})
    assert response.status_code == 400
    assert "Embeddings" in response.json()["detail"]


# ---- SharePoint MCP ----


def test_sharepoint_mcp_mock_list_and_call():
    tools = client.post("/connectors/sharepoint-mcp", json={"operation": "listTools", "live": False}).json()["tools"]
    assert tools and all("name" in t for t in tools)
    called = client.post("/connectors/sharepoint-mcp", json={"tool": "search_documents", "arguments": {"query": "x"}, "live": False}).json()
    assert called["tool"] == "search_documents" and called["structured"]["documents"]


def test_sharepoint_mcp_live_call_with_fake_session():
    class FakeSession:
        async def call_tool(self, name, arguments):
            text = SimpleNamespace(type="text", text=f"{name}:{arguments['q']}")
            return SimpleNamespace(content=[text], isError=False, structuredContent={"n": 1})

        async def list_tools(self):
            return SimpleNamespace(tools=[SimpleNamespace(name="t", description=None, inputSchema={})])

    @contextlib.asynccontextmanager
    async def factory():
        yield FakeSession()

    tool = SharePointMCPTool(session_factory=factory)
    assert asyncio.run(tool.call_tool("search", {"q": "x"})) == {"tool": "search", "text": "search:x", "structured": {"n": 1}}
    assert asyncio.run(tool.list_tools())["tools"][0] == {"name": "t", "description": "", "inputSchema": {}}


def test_sharepoint_mcp_tool_error_is_raised():
    class FailingSession:
        async def call_tool(self, name, arguments):
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="denied")], isError=True, structuredContent=None)

    @contextlib.asynccontextmanager
    async def factory():
        yield FailingSession()

    try:
        asyncio.run(SharePointMCPTool(session_factory=factory).call_tool("search", {}))
    except SharePointMCPError as err:
        assert "denied" in str(err) and err.status_code == 502
    else:
        raise AssertionError("expected SharePointMCPError")


# ---- Parser ----


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def test_parse_text_docx_xlsx_and_image():
    from docx import Document
    from openpyxl import Workbook

    assert client.post("/tools/parse", json={"name": "a.txt", "content": "hello"}).json()["text"] == "hello"

    doc = Document()
    doc.add_paragraph("Material Register")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "UPS", "Delivered"
    buf = io.BytesIO()
    doc.save(buf)
    text = client.post("/tools/parse", json={"name": "r.docx", "content": _b64(buf.getvalue()), "encoding": "base64"}).json()["text"]
    assert "Material Register" in text and "UPS | Delivered" in text

    wb = Workbook()
    wb.active.append(["Item", "Qty"])
    wb.active.append(["CAT6", 12])
    buf = io.BytesIO()
    wb.save(buf)
    text = client.post("/tools/parse", json={"name": "b.xlsx", "content": _b64(buf.getvalue()), "encoding": "base64"}).json()["text"]
    assert "CAT6 | 12" in text

    image = client.post("/tools/parse", json={"name": "scan.png", "content": _b64(b"\x89PNG"), "encoding": "base64"}).json()
    assert image["text"] == "" and "OCR" in image["unsupported"]


def test_parse_corrupt_pdf_is_422():
    response = client.post("/tools/parse", json={"name": "x.pdf", "content": _b64(b"not a pdf"), "encoding": "base64"})
    assert response.status_code == 422


# ---- SharePoint batch read ----


def test_sharepoint_get_folder_files_mock():
    body = client.post("/connectors/sharepoint", json={"operation": "getFolderFiles", "path": "Projects/P1", "live": False}).json()
    assert len(body["files"]) == 2 and all(f["content"] for f in body["files"])
    assert body["files"][0]["path"].startswith("/Projects/P1/")


def test_gemini_json_mode_does_not_send_json_object(monkeypatch):
    """Gemini's OpenAI-compatible API is called without response_format; the reply is parsed from text."""
    import openai

    calls = []

    class FakeCompletions:
        async def create(self, **kwargs):
            calls.append(kwargs)
            message = SimpleNamespace(content='```json\n{"blocks": []}\n```')
            return SimpleNamespace(choices=[SimpleNamespace(message=message)])

    class FakeClient:
        def __init__(self, **kwargs):
            calls.append(kwargs)
            self.chat = SimpleNamespace(completions=FakeCompletions())

    monkeypatch.setattr(openai, "AsyncOpenAI", FakeClient)
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    response = client.post("/agent/execute", json={"model": "gemini:gemini-3.8-flash", "prompt": "Hi", "response_format": "json"})
    assert response.status_code == 200, response.text
    assert response.json()["json"] == {"blocks": []}
    assert calls[0]["base_url"].startswith("https://generativelanguage.googleapis.com/")
    assert "response_format" not in calls[1]


def test_qdrant_upsert_creates_collection_and_indexes():
    class FakeClient:
        def __init__(self):
            self.calls = []

        async def collection_exists(self, name):
            return False

        async def create_collection(self, **kwargs):
            self.calls.append(("create_collection", kwargs["vectors_config"].size))

        async def create_payload_index(self, **kwargs):
            self.calls.append(("index", kwargs["field_name"]))

        async def upsert(self, **kwargs):
            self.calls.append(("upsert", len(kwargs["points"])))

    fake = FakeClient()
    points = [{"vector": [0.1, 0.2, 0.3], "payload": {"url": "u", "chunk_index": 0, "project_id": "P1"}}]
    result = asyncio.run(QdrantTool(client=fake).upsert("docs", points, index_fields=["project_id", "doc_type"]))
    assert fake.calls == [("create_collection", 3), ("index", "project_id"), ("index", "doc_type"), ("upsert", 1)]
    assert result == {"collection": "docs", "upserted": 1, "indexed_fields": ["project_id", "doc_type"]}
