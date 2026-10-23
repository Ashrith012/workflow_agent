"""Extracts plain text from documents fetched by the SharePoint connector."""
import base64
import io
from typing import Any, Dict

TEXT_EXTENSIONS = {"txt", "md", "csv", "json", "xml", "html", "htm", "log"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp"}


class ParserError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def parse_document(name: str, content: str, encoding: str = "utf-8") -> Dict[str, Any]:
    """`content` is text when `encoding` is utf-8, else base64 (the SharePoint getFile shape)."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    raw = content.encode("utf-8") if encoding == "utf-8" else base64.b64decode(content)

    if ext in IMAGE_EXTENSIONS:
        # Scanned documents and images need OCR, which this parser does not do.
        return {"name": name, "text": "", "unsupported": "image (OCR not available)"}
    if ext == "pdf":
        return {"name": name, **_pdf(raw)}
    if ext == "docx":
        return {"name": name, "text": _docx(raw)}
    if ext in ("xlsx", "xlsm"):
        return {"name": name, "text": _xlsx(raw)}
    if ext in TEXT_EXTENSIONS or encoding == "utf-8":
        return {"name": name, "text": raw.decode("utf-8", errors="replace")}
    return {"name": name, "text": "", "unsupported": f".{ext or '?'} files"}


def _pdf(raw: bytes) -> Dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages) if t.strip())
    result: Dict[str, Any] = {"text": text, "pages": len(pages)}
    if pages and not text.strip():
        result["unsupported"] = "scanned PDF (no text layer; OCR not available)"
    return result


def _docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                parts.append(" | ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(parts)
